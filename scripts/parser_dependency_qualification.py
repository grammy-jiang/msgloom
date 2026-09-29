#!/usr/bin/env python3
"""Exercise parser dependency candidate APIs against synthetic fixtures."""

from __future__ import annotations

import json
import platform
from email import policy
from email.parser import BytesParser
from importlib.metadata import version
from io import BytesIO

import openpyxl.xml
import pypdfium2 as pdfium
from docx import Document
from openpyxl import load_workbook
from python_calamine import CalamineWorkbook
from selectolax.lexbor import LexborHTMLParser

from tests.parser_fixture_builders import (
    HTML_BODY,
    build_docx_fixture,
    build_mime_fixture,
    build_pdf_fixture,
    build_xlsx_fixture,
)

CANDIDATES = {
    "defusedxml": "0.7.1",
    "lxml": "6.1.3",
    "openpyxl": "3.1.5",
    "pypdfium2": "5.13.0",
    "python-calamine": "0.8.2",
    "python-docx": "1.2.0",
    "selectolax": "0.4.12",
}


def _require(condition: bool, detail: str) -> None:
    if not condition:
        raise RuntimeError(detail)


def _exercise_mime() -> dict[str, object]:
    message = BytesParser(policy=policy.default).parsebytes(build_mime_fixture())
    content_types = [part.get_content_type() for part in message.walk()]
    html_part = next(
        part for part in message.walk() if part.get_content_type() == "text/html"
    )
    html = html_part.get_content()
    _require("Alpha" in html and "Tail." in html, "MIME HTML decode lost content")
    return {"content_types": content_types, "html_decoded": True}


def _exercise_html() -> dict[str, object]:
    document = LexborHTMLParser(HTML_BODY)
    link = document.css_first("a")
    table = document.css_first("table")
    _require(link is not None and table is not None, "HTML fixture nodes missing")
    href = link.attributes.get("href")
    _require(href == "https://example.invalid/reference", "HTML link changed")
    _require("First" in table.text() and "7" in table.text(), "HTML table lost data")
    return {"link": href, "table_text_observed": True}


def _exercise_excel() -> dict[str, object]:
    fixture = build_xlsx_fixture()
    workbook = CalamineWorkbook.from_filelike(BytesIO(fixture))
    sheet = workbook.get_sheet_by_name("Visible")
    metadata = workbook.sheets_metadata
    _require(sheet.start == (2, 1), "Calamine source start offset changed")
    _require(
        sheet.merged_cell_ranges == [((2, 3), (2, 4))],
        "Calamine merged range changed",
    )
    _require(
        any(
            item.name == "Hidden" and str(item.visible).endswith(".Hidden")
            for item in metadata
        ),
        "Calamine hidden-sheet metadata changed",
    )
    rows = sheet.to_python()
    _require(rows[0][0] == 1.0 and rows[1][1] == 3.0, "Calamine values changed")

    formula_book = load_workbook(BytesIO(fixture), data_only=False)
    cached_book = load_workbook(BytesIO(fixture), data_only=True)
    formula_sheet = formula_book["Visible"]
    cached_sheet = cached_book["Visible"]
    _require(formula_sheet["C4"].value == "=SUM(B3,2)", "Formula text changed")
    _require(cached_sheet["C4"].value == 3, "Cached formula value changed")
    _require(
        str(next(iter(formula_sheet.merged_cells.ranges))) == "D3:E3",
        "OpenPyXL merged range changed",
    )
    _require(formula_book["Hidden"].sheet_state == "hidden", "Hidden state changed")
    _require(openpyxl.xml.DEFUSEDXML is True, "OpenPyXL defusedxml is not active")
    return {
        "calamine_start": list(sheet.start),
        "cached_formula_value": cached_sheet["C4"].value,
        "defusedxml": openpyxl.xml.DEFUSEDXML,
        "formula": formula_sheet["C4"].value,
        "merge": "D3:E3",
    }


def _exercise_pdf() -> dict[str, object]:
    document = pdfium.PdfDocument(build_pdf_fixture())
    _require(len(document) == 2, "PDF fixture page count changed")
    first = document[0].get_textpage().get_text_range()
    second = document[1].get_textpage().get_text_range()
    _require("First line" in first and "Second line" in first, "PDF text changed")
    _require(not second.strip(), "Image-only PDF page unexpectedly has text")
    return {"image_only_page_has_text": False, "page_count": len(document)}


def _exercise_docx() -> dict[str, object]:
    document = Document(BytesIO(build_docx_fixture()))
    inner_types = [type(item).__name__ for item in document.iter_inner_content()]
    _require(
        inner_types[:3] == ["Paragraph", "Table", "Paragraph"],
        "DOCX paragraph/table order changed",
    )
    ole_relationships = [
        rel for rel in document.part.rels.values() if rel.reltype.endswith("/oleObject")
    ]
    _require(len(ole_relationships) == 1, "DOCX OLE relationship changed")
    return {"inner_types": inner_types, "ole_relationships": len(ole_relationships)}


def main() -> None:
    """Run candidate API checks and emit deterministic JSON evidence."""
    installed = {name: version(name) for name in CANDIDATES}
    _require(installed == CANDIDATES, "Installed candidate versions differ")
    result = {
        "candidate_versions": installed,
        "checks": {
            "docx": _exercise_docx(),
            "excel": _exercise_excel(),
            "html": _exercise_html(),
            "mime": _exercise_mime(),
            "pdf": _exercise_pdf(),
        },
        "host": {
            "machine": platform.machine(),
            "python": platform.python_version(),
            "python_implementation": platform.python_implementation(),
            "system": platform.system(),
        },
        "qualification": f"runtime-exercised-cp{platform.python_version_tuple()[0]}{platform.python_version_tuple()[1]}-synthetic-only",
    }
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
