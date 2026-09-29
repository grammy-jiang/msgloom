"""Exercise production parser dispatch through the actual Linux boundary."""

from __future__ import annotations

import asyncio
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document

from msgloom.preparation.contracts import (
    DocumentFormat,
    ParserConfig,
    ParserLimits,
    ParserProvenance,
    ParserRequest,
    SavedByteReference,
    TableBlock,
    TextBlock,
)
from msgloom.preparation.isolation import parse_isolated
from msgloom.preparation.isolation.registry import PRODUCTION_REGISTRY
from tests.parser_pdf.fixtures import build_text_pdf


def _docx_content() -> bytes:
    document = Document()
    document.add_paragraph("Synthetic boundary 42")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Mapped cell 42"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _request(content: bytes, format_: DocumentFormat, profile: str) -> ParserRequest:
    return ParserRequest(
        source=SavedByteReference(
            reference=f"synthetic-boundary:{format_.value}",
            sha256=sha256(content).hexdigest(),
            byte_count=len(content),
        ),
        detected_format=format_,
        parser=PRODUCTION_REGISTRY.resolve(format_).identity,
        config=ParserConfig(profile=profile),
        limits=ParserLimits(
            wall_time_seconds=20,
            memory_bytes=512 * 1024 * 1024,
            decompressed_bytes=4 * 1024 * 1024,
            output_bytes=1024 * 1024,
            container_members=128,
        ),
    )


@pytest.mark.parametrize(
    ("format_", "profile", "content"),
    [
        (DocumentFormat.TEXT, "mime-html-v1", b"Synthetic boundary 42"),
        (
            DocumentFormat.HTML,
            "mime-html-v1",
            b"<p>Synthetic boundary 42</p>",
        ),
        (
            DocumentFormat.JSON,
            "mime-html-v1",
            b'{"value":"Synthetic boundary 42"}',
        ),
        (
            DocumentFormat.MIME,
            "mime-html-v1",
            (
                b"MIME-Version: 1.0\r\nContent-Type: text/plain; charset=utf-8"
                b"\r\n\r\nSynthetic boundary 42\r\n"
            ),
        ),
        (
            DocumentFormat.PDF,
            "pdf-primary-v1",
            build_text_pdf(b"Synthetic boundary 42"),
        ),
        (DocumentFormat.DOCX, "word-native-v1", _docx_content()),
    ],
)
def test_production_parser_preserves_content_and_provenance_in_sandbox(
    format_: DocumentFormat,
    profile: str,
    content: bytes,
) -> None:
    """Real native dependencies must load inside each trusted runtime mount."""
    request = _request(content, format_, profile)
    output = asyncio.run(parse_isolated(request, content))
    if output.provenance != ParserProvenance.from_request(request):
        pytest.fail("isolated format parser changed source or configuration lineage")
    text = "\n".join(
        block.text for block in output.blocks if isinstance(block, TextBlock)
    )
    if "Synthetic boundary 42" not in text:
        pytest.fail("isolated production parser lost the synthetic source text")
    if format_ is DocumentFormat.DOCX:
        cells = [
            cell
            for block in output.blocks
            if isinstance(block, TableBlock)
            for cell in block.table.cells
        ]
        if len(cells) != 1 or cells[0].text != "Mapped cell 42":
            pytest.fail("isolated Word parser lost the mapped table cell")
