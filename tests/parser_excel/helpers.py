"""Synthetic spreadsheet fixtures and request builders for the Excel lane."""

from __future__ import annotations

from datetime import date, time
from hashlib import sha256
from io import BytesIO
from typing import cast
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from msgloom.preparation.contracts import (
    DocumentFormat,
    ParserConfig,
    ParserIdentity,
    ParserLimits,
    ParserRequest,
    SavedByteReference,
)
from msgloom.preparation.parsers.excel import (
    BACKEND,
    PARSER_NAME,
    PARSER_VERSION,
)


def request_for(
    content: bytes,
    fmt: DocumentFormat,
    *,
    settings: tuple[tuple[str, str], ...] = (),
    decompressed_bytes: int = 2_000_000,
    output_bytes: int = 2_000_000,
    container_members: int = 200,
) -> ParserRequest:
    """Build a request whose source identity exactly matches fixture bytes."""
    return ParserRequest(
        source=SavedByteReference(
            reference=f"fixture:{fmt.value}",
            sha256=sha256(content).hexdigest(),
            byte_count=len(content),
        ),
        detected_format=fmt,
        parser=ParserIdentity(
            name=PARSER_NAME,
            version=PARSER_VERSION,
            backend=BACKEND,
        ),
        config=ParserConfig(profile="excel-primary-v1", settings=settings),
        limits=ParserLimits(
            wall_time_seconds=5,
            memory_bytes=64 * 1024 * 1024,
            decompressed_bytes=decompressed_bytes,
            output_bytes=output_bytes,
            container_members=container_members,
        ),
    )


def build_visibility_fixture() -> bytes:
    """Build OOXML with hidden content, link, error, merge, and formula."""
    workbook = Workbook()
    visible = cast(Worksheet, workbook.active)
    visible.title = "Visible"
    visible["B3"] = "shown"
    visible["C3"] = "#DIV/0!"
    visible["D3"] = "hidden-column"
    visible["E3"] = "linked"
    visible["E3"].hyperlink = "https://example.invalid/workbook"
    visible["B4"] = "hidden-row"
    visible["C4"] = "=SUM(1,2)"
    visible["B5"] = date(2026, 9, 29)
    visible["C5"] = time(10, 11, 12)
    visible["B5"].number_format = "yyyy-mm-dd"
    visible["C5"].number_format = "hh:mm:ss"
    visible.merge_cells("F3:G3")
    visible["F3"] = "merged"
    visible.row_dimensions[4].hidden = True
    visible.column_dimensions["D"].hidden = True

    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "secret"
    hidden.sheet_state = "hidden"

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def add_unsupported_ooxml_parts(content: bytes) -> bytes:
    """Add inert synthetic members used only for limitation detection."""
    source = ZipFile(BytesIO(content))
    buffer = BytesIO()
    try:
        with ZipFile(buffer, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"application/vnd.openxmlformats-officedocument."
                        b"spreadsheetml.sheet.main+xml",
                        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                    )
                target.writestr(info, data)
            target.writestr("xl/vbaProject.bin", b"synthetic-vba")
            target.writestr("xl/embeddings/oleObject1.bin", b"synthetic-object")
            target.writestr(
                "xl/externalLinks/externalLink1.xml",
                b"<externalLink/>",
            )
    finally:
        source.close()
    return buffer.getvalue()


def build_ods_fixture() -> bytes:
    """Build a small ODS with leading blanks and several typed values."""
    content = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b"<office:document-content "
        b'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        b'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        b'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        b'office:version="1.2"><office:body><office:spreadsheet>'
        b'<table:table table:name="Data"><table:table-row>'
        b"<table:table-cell/><table:table-cell/></table:table-row>"
        b"<table:table-row><table:table-cell/>"
        b'<table:table-cell office:value-type="string">'
        b"<text:p>ods-value</text:p></table:table-cell>"
        b'<table:table-cell office:value-type="float" office:value="2.5"/>'
        b'<table:table-cell office:value-type="boolean" '
        b'office:boolean-value="true"/>'
        b'<table:table-cell office:value-type="date" '
        b'office:date-value="2026-09-29"/>'
        b'<table:table-cell office:value-type="time" '
        b'office:time-value="PT10H11M12S"/>'
        b"</table:table-row></table:table></office:spreadsheet>"
        b"</office:body></office:document-content>"
    )
    manifest = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b"<manifest:manifest "
        b'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        b'manifest:version="1.2"><manifest:file-entry manifest:full-path="/" '
        b'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        b'<manifest:file-entry manifest:full-path="content.xml" '
        b'manifest:media-type="text/xml"/></manifest:manifest>'
    )
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.spreadsheet",
            compress_type=ZIP_STORED,
        )
        archive.writestr(
            "content.xml",
            content,
            compress_type=ZIP_DEFLATED,
        )
        archive.writestr(
            "META-INF/manifest.xml",
            manifest,
            compress_type=ZIP_DEFLATED,
        )
    return buffer.getvalue()
