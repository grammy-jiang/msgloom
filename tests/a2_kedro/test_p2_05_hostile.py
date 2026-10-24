"""P2-05 hostile workbooks: huge dimension claims, bombs, DTDs, time."""

from __future__ import annotations

import time
import zipfile
from io import BytesIO

import openpyxl
import pytest

from msgloom.a2_kedro.parsers import sheet_scan as ss
from msgloom.a2_kedro.parsers import spreadsheets as sp
from msgloom.a2_kedro.parsers.spreadsheets import Limitation as L
from msgloom.a2_kedro.parsers.spreadsheets import Outcome, parse_spreadsheet
from tests.a2_kedro.helpers import check

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
#: Generous absolute wall-clock bound. It only catches runaway work.
SLOW = 30.0
BOOK = "application/vnd.openxmlformats-officedocument.spreadsheetml"


def mini(sheet_xml: str, extra: dict[str, bytes] | None = None) -> bytes:
    """Build a tiny workbook by hand with one sheet body."""
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types"><Override PartName="/xl/workbook.xml" '
            f'ContentType="{BOOK}.sheet.main+xml"/><Override '
            f'PartName="/xl/worksheets/sheet1.xml" ContentType="{BOOK}.worksheet+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            f'<Relationships xmlns="{PKG}"><Relationship Id="r1" '
            f'Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets>'
            '<sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" '
            f'Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>'
        ),
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{MAIN}">{sheet_xml}</worksheet>',
    }
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in parts.items():
            zf.writestr(name, body)
        for name, body in (extra or {}).items():
            zf.writestr(name, body)
    return out.getvalue()


def inline(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def test_hand_built_workbook_parses() -> None:
    body = f'<sheetData><row r="1">{inline("A1", "hi")}</row></sheetData>'
    result = parse_spreadsheet(mini(body))
    check(result.outcome is Outcome.COMPLETE, "complete")
    check(result.cells[0].value == "hi" and result.cells[0].ref == "A1", "cell")


def test_declared_full_grid_dimension_does_not_allocate() -> None:
    body = (
        '<dimension ref="A1:XFD1048576"/><sheetData><row r="1">'
        f"{inline('A1', 'x')}</row></sheetData>"
    )
    begin = time.perf_counter()
    result = parse_spreadsheet(mini(body))
    check(time.perf_counter() - begin < SLOW, "declared dimension cost time")
    check(result.outcome is Outcome.COMPLETE and len(result.cells) == 1, "ok")
    check((result.sheets[0].row_count, result.sheets[0].column_count) == (1, 1), "dims")


def test_far_sparse_cell_is_bounded_and_flagged() -> None:
    body = (
        '<dimension ref="A1:XFD1048576"/><sheetData>'
        f'<row r="1">{inline("A1", "x")}</row>'
        f'<row r="1048576">{inline("XFD1048576", "far")}</row></sheetData>'
    )
    begin = time.perf_counter()
    result = parse_spreadsheet(mini(body))
    check(time.perf_counter() - begin < SLOW, "far cell cost too much time")
    check(result.outcome is Outcome.INCOMPLETE, "partial")
    check(L.DIMENSION_LIMIT in result.limitations, "dimension limit")
    check(result.cells[0].block_id == "s0!A1", "near cell kept")


def test_styled_empty_cells_do_not_inflate_extents() -> None:
    row = "".join(f'<c r="{c}5" s="1"/>' for c in ("A", "B", "XFD"))
    body = f'<sheetData><row r="5">{row}</row></sheetData>'
    result = parse_spreadsheet(mini(body))
    check(result.outcome is Outcome.COMPLETE and not result.cells, "empty styled row")


@pytest.mark.parametrize(
    "prefix",
    [
        '<!DOCTYPE x [<!ENTITY a "aaaa">]>',
        '<!DOCTYPE x SYSTEM "http://example.invalid/x.dtd">',
    ],
)
def test_doctype_and_entities_are_rejected(prefix) -> None:
    data = mini("<sheetData/>")
    with zipfile.ZipFile(BytesIO(data)) as zf:
        files = {n: zf.read(n) for n in zf.namelist()}
    files["xl/worksheets/sheet1.xml"] = (
        prefix.encode() + files["xl/worksheets/sheet1.xml"]
    )
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        for name, body in files.items():
            zf.writestr(name, body)
    result = parse_spreadsheet(out.getvalue())
    check(result.outcome is Outcome.FAILED, "failed")
    check(result.limitations == (L.CORRUPT,), "corrupt")


def test_zip_bomb_stops_at_uncompressed_bound() -> None:
    bomb = {"xl/media/big.bin": b"\x00" * (80 * 1024 * 1024)}
    data = mini("<sheetData/>", bomb)
    check(len(data) < 1024 * 1024, "fixture must be small")
    begin = time.perf_counter()
    result = parse_spreadsheet(data)
    check(time.perf_counter() - begin < SLOW, "bomb took too long")
    check(result.outcome is Outcome.INCOMPLETE, "partial")
    check(L.ZIP_LIMIT in result.limitations and not result.cells, "zip limit")


def test_declared_size_smaller_than_actual_is_corrupt() -> None:
    data = bytearray(mini("<sheetData/>", {"xl/media/lie.bin": b"A" * 5000}))
    for sig, off in ((b"PK\x03\x04", 22), (b"PK\x01\x02", 24)):
        pos = data.find(sig)
        while pos >= 0:
            if b"lie.bin" in bytes(data[pos : pos + 80]):
                data[pos + off : pos + off + 4] = (10).to_bytes(4, "little")
            pos = data.find(sig, pos + 4)
    result = parse_spreadsheet(bytes(data))
    check(L.CORRUPT in result.limitations, "lying size is corrupt")


def test_many_cells_parse_in_linear_time() -> None:
    def book(count: int) -> bytes:
        rows = "".join(
            f'<row r="{i}"><c r="A{i}"><v>{i}</v></c></row>' for i in range(1, count)
        )
        return mini(f"<sheetData>{rows}</sheetData>")

    small, large = book(2000), book(8000)
    t0 = time.perf_counter()
    parse_spreadsheet(small)
    t1 = time.perf_counter()
    parse_spreadsheet(large)
    t2 = time.perf_counter()
    check(t2 - t1 < SLOW, "8000 rows took too long")
    check((t2 - t1) < max(t1 - t0, 0.05) * 30, "time grows faster than linear")


def test_failure_messages_do_not_echo_input() -> None:
    marker = b"private-marker-9917"
    result = parse_spreadsheet(b"PK\x03\x04" + marker * 4)
    check(marker.decode() not in repr(result), "repr echoes input")
    check(all(marker.decode() not in i.value for i in result.limitations), "limits")


# --- parts that openpyxl loads in full ------------------------------------

SST = f"{BOOK}.sharedStrings+xml"
#: Modest sizes just above the budgets (old code needs seconds, not minutes).
STYLE_N, STRING_N, PART_N = 60_000, 120_000, 60_000


def spy_load(monkeypatch) -> list[tuple[bytes, dict]]:
    """Record (package bytes, keyword arguments) of each openpyxl load."""
    real = sp.openpyxl.load_workbook
    calls: list[tuple[bytes, dict]] = []

    def spy(source, **kwargs):
        calls.append((source.getvalue(), kwargs))
        return real(source, **kwargs)

    monkeypatch.setattr(sp.openpyxl, "load_workbook", spy)
    return calls


def with_parts(parts: dict[str, bytes], types: str = "") -> bytes:
    """Return a one-cell workbook plus ``parts`` and extra content types."""
    body = f'<sheetData><row r="1">{inline("A1", "hi")}</row></sheetData>'
    data = mini(body, parts)
    if types:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            files = {n: zf.read(n) for n in zf.namelist()}
        key = "[Content_Types].xml"
        files[key] = files[key].replace(b"</Types>", types.encode() + b"</Types>")
        out = BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, blob in files.items():
                zf.writestr(name, blob)
        data = out.getvalue()
    return data


def styles(count: int) -> bytes:
    head = f'<styleSheet xmlns="{MAIN}"><cellXfs count="{count}">'.encode()
    return head + b"<xf/>" * count + b"</cellXfs></styleSheet>"


def strings(count: int, text: bytes = b"a") -> bytes:
    head = f'<sst xmlns="{MAIN}" count="{count}" uniqueCount="{count}">'.encode()
    return head + (b"<si><t>" + text + b"</t></si>") * count + b"</sst>"


def override(part: str) -> str:
    return f'<Override PartName="/{part}" ContentType="{SST}"/>'


def refused(data: bytes, monkeypatch, why: str) -> None:
    calls = spy_load(monkeypatch)
    begin = time.perf_counter()
    result = parse_spreadsheet(data)
    check(time.perf_counter() - begin < SLOW, f"{why}: took too long")
    check(not calls, f"{why}: openpyxl loaded the workbook")
    check(result.outcome is Outcome.INCOMPLETE, f"{why}: outcome")
    check(L.SCAN_LIMIT in result.limitations, f"{why}: limitation")
    check(not result.cells and not result.sheets, f"{why}: content returned")
    check(result.to_receipt("0") is not None, f"{why}: receipt")


def test_oversized_styles_part_is_refused_before_openpyxl(monkeypatch) -> None:
    data = with_parts({"xl/styles.xml": styles(STYLE_N)})
    refused(data, monkeypatch, "styles")


@pytest.mark.parametrize("part", ["xl/sharedStrings.xml", "xl/odd/names.xml"])
def test_oversized_shared_strings_are_refused_before_openpyxl(
    monkeypatch, part
) -> None:
    # The second part name is found by content type only, as openpyxl does.
    data = with_parts({part: strings(STRING_N)}, override(part))
    refused(data, monkeypatch, "shared strings")


def test_one_huge_shared_string_is_refused_by_byte_bound(monkeypatch) -> None:
    text = b"a" * (ss.MAX_PART_BYTES + 1)
    refused(
        with_parts({"xl/sharedStrings.xml": strings(1, text)}), monkeypatch, "bytes"
    )


@pytest.mark.parametrize(
    "part", ["docProps/core.xml", "xl/worksheets/_rels/sheet1.xml.rels"]
)
def test_other_loaded_parts_have_an_element_bound(monkeypatch, part) -> None:
    body = b"<r>" + b"<x/>" * PART_N + b"</r>"
    refused(with_parts({part: body}), monkeypatch, "other part")


def test_part_budgets_are_named_and_below_hostile_sizes() -> None:
    check(ss.MAX_STYLE_ELEMENTS < STYLE_N, "style bound")
    check(ss.MAX_SHARED_STRINGS < STRING_N, "string bound")
    check(ss.MAX_STRING_ELEMENTS >= ss.MAX_SHARED_STRINGS, "string elements")
    check(ss.MAX_PART_ELEMENTS < PART_N, "part bound")


def test_ordinary_styles_and_strings_still_complete(monkeypatch) -> None:
    calls = spy_load(monkeypatch)
    book = openpyxl.Workbook()
    sheet = book.create_sheet("Data")
    for i in range(1, 1001):
        sheet.cell(i, 1, f"text {i}")
        sheet.cell(i, 2, i).number_format = f"0.{'0' * (i % 7)}"
    out = BytesIO()
    book.save(out)
    result = parse_spreadsheet(out.getvalue())
    check(result.outcome is Outcome.COMPLETE, "ordinary parts complete")
    check(len(calls) == 2 and len(result.cells) == 2000, "loaded and read")


# --- mutation gaps: column clip and capped sheet stubs -------------------


def test_hidden_column_range_is_clipped_to_the_last_column() -> None:
    cols = '<cols><col min="1" max="1000000" hidden="1"/></cols>'
    body = f'{cols}<sheetData><row r="1">{inline("A1", "x")}</row></sheetData>'
    result = parse_spreadsheet(mini(body))
    check(result.outcome is Outcome.COMPLETE, "clipped range stays in budget")
    check(result.sheets[0].hidden_columns == tuple(range(16_384)), "clip")


SHEET = "xl/worksheets/sheet1.xml"


def capped_book(count: int) -> bytes:
    rows = "".join(
        f'<row r="{i}"><c r="A{i}"><v>{i}</v></c></row>' for i in range(1, count + 1)
    )
    return mini(f"<sheetData>{rows}</sheetData>")


def test_capped_sheet_gives_no_cells_and_openpyxl_never_parses_it(
    monkeypatch,
) -> None:
    monkeypatch.setattr(sp, "Budget", lambda: ss.Budget(elements=40))
    calls = spy_load(monkeypatch)
    result = parse_spreadsheet(capped_book(100))
    check(result.outcome is Outcome.INCOMPLETE, "partial")
    check(L.SCAN_LIMIT in result.limitations and not result.cells, "no cells")
    check(len(calls) == 2, "workbook still opened for the sheet list")
    for sent, _kwargs in calls:  # both passes read the stub, not the sheet
        with zipfile.ZipFile(BytesIO(sent)) as zf:
            part = zf.read(SHEET)
        check(b"<c " not in part and b"<row" not in part, "sheet part not stubbed")


def test_drop_capped_stubs_only_capped_sheet_parts() -> None:
    with zipfile.ZipFile(BytesIO(capped_book(3))) as zf:
        check(ss.drop_capped(zf, {SHEET: ss.SheetScan()}) is None, "nothing capped")
        capped = ss.SheetScan(capped=True)
        out = ss.drop_capped(zf, {SHEET: capped})
        check(out is not None, "capped sheet is replaced")
        with zipfile.ZipFile(BytesIO(out or b"")) as new:
            check(new.read(SHEET) == ss._STUB, "stub content")
            check(sorted(new.namelist()) == sorted(zf.namelist()), "members kept")
            check(new.read("xl/workbook.xml") == zf.read("xl/workbook.xml"), "kept")
