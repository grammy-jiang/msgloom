"""P2-05 spreadsheet parser: sheets, cells, formulas, hidden state, limits."""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from io import BytesIO
from typing import Any

import openpyxl
import pytest

from msgloom.a2_kedro.contracts.codec import A2ContractError
from msgloom.a2_kedro.parsers import spreadsheets as sp
from msgloom.a2_kedro.parsers.spreadsheets import (
    Limitation as L,
)
from msgloom.a2_kedro.parsers.spreadsheets import Outcome, parse_spreadsheet
from tests.a2_kedro.helpers import check

SECRET = "tiny-secret-value-7431"


def active(workbook: openpyxl.Workbook) -> Any:
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("workbook has no active sheet")
    return sheet


def save(workbook: openpyxl.Workbook) -> bytes:
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def rezip(data: bytes, edit=None, extra: dict[str, bytes] | None = None) -> bytes:
    """Copy a package; ``edit(name, bytes)`` may change a member."""
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            body = src.read(info)
            dst.writestr(info.filename, edit(info.filename, body) if edit else body)
        for name, body in (extra or {}).items():
            dst.writestr(name, body)
    return out.getvalue()


def simple() -> bytes:
    wb = openpyxl.Workbook()
    ws = active(wb)
    ws.title = "Data"
    ws["A1"] = SECRET
    ws["B1"] = 3.5
    ws["C1"] = True
    ws["A2"] = dt.datetime(2026, 1, 2, 3, 4, 5)
    ws["B3"] = "#DIV/0!"
    return save(wb)


def cells_by_ref(result):
    return {c.block_id: c for c in result.cells}


def test_simple_workbook_is_complete_with_stable_block_ids() -> None:
    result = parse_spreadsheet(simple())
    check(result.outcome is Outcome.COMPLETE and not result.limitations, "complete")
    check(result.format == "xlsx", "format")
    check(result.sheets[0].block_id == "s0", "sheet id")
    got = cells_by_ref(result)
    check(sorted(got) == ["s0!A1", "s0!A2", "s0!B1", "s0!B3", "s0!C1"], "ids")
    kinds = {ref: c.kind for ref, c in got.items()}
    check(kinds["s0!A1"] == "text" and kinds["s0!B1"] == "number", "kinds")
    check(kinds["s0!C1"] == "boolean" and kinds["s0!A2"] == "date", "kinds 2")
    check(got["s0!A1"].value == SECRET, "value")
    check((result.sheets[0].row_count, result.sheets[0].column_count) == (3, 3), "dims")


def test_error_value_has_kind_error() -> None:
    wb = openpyxl.Workbook()
    ws = active(wb)
    ws["A1"] = "=1/0"
    data = rezip(
        save(wb),
        lambda n, b: (
            b.replace(b"<v></v>", b"<v>#DIV/0!</v>").replace(
                b'<c r="A1"', b'<c r="A1" t="e"'
            )
            if n.endswith("sheet1.xml")
            else b
        ),
    )
    cell = parse_spreadsheet(data).cells[0]
    check(cell.kind == "error" and cell.formula == "=1/0", "error cell")


def formula_book() -> bytes:
    wb = openpyxl.Workbook()
    active(wb)["A1"] = 1
    active(wb)["A2"] = "=A1+2"
    return save(wb)


def test_formula_without_cached_value_is_flagged_not_empty() -> None:
    result = parse_spreadsheet(formula_book())
    cell = cells_by_ref(result)["s0!A2"]
    check(cell.formula == "=A1+2" and cell.value is None, "formula kept")
    check(cell.formula_uncached, "uncached flag")
    check(result.outcome is Outcome.INCOMPLETE, "never complete")
    check(L.FORMULA_WITHOUT_CACHE in result.limitations, "limitation")


def with_cache(data: bytes, cached: str) -> bytes:
    def edit(name: str, body: bytes) -> bytes:
        if not name.endswith("sheet1.xml"):
            return body
        return re.sub(rb"<v\s*/>|<v></v>", f"<v>{cached}</v>".encode(), body)

    return rezip(data, edit)


def test_formula_text_and_cached_value_stay_distinct() -> None:
    result = parse_spreadsheet(with_cache(formula_book(), "3"))
    cell = cells_by_ref(result)["s0!A2"]
    check(cell.formula == "=A1+2" and cell.value == 3, "both kept")
    check(not cell.formula_uncached and cell.kind == "number", "cached")
    check(result.outcome is Outcome.COMPLETE, "cached formula is complete")


def test_hidden_sheets_rows_and_columns_are_flagged() -> None:
    wb = openpyxl.Workbook()
    ws = active(wb)
    ws["A1"], ws["B2"], ws["C3"] = "a", "b", "c"
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["C"].hidden = True
    wb.create_sheet("H1").sheet_state = "hidden"
    wb.create_sheet("H2").sheet_state = "veryHidden"
    result = parse_spreadsheet(save(wb))
    check(
        [s.state for s in result.sheets] == ["visible", "hidden", "veryHidden"],
        "states",
    )
    check(result.sheets[0].hidden_rows == (1,), "hidden rows")
    check(result.sheets[0].hidden_columns == (2,), "hidden cols")
    got = cells_by_ref(result)
    check(got["s0!B2"].hidden_row and not got["s0!A1"].hidden_row, "row flag")
    check(got["s0!C3"].hidden_column and not got["s0!A1"].hidden_column, "col flag")


def test_merged_cells_are_unsupported_and_partial() -> None:
    wb = openpyxl.Workbook()
    active(wb)["A1"] = "x"
    active(wb).merge_cells("A1:B1")
    result = parse_spreadsheet(save(wb))
    check(result.outcome is Outcome.INCOMPLETE, "partial")
    check(L.MERGED_CELLS in result.limitations, "merged limitation")


FEATURES = {
    "xl/comments1.xml": L.COMMENTS,
    "xl/charts/chart1.xml": L.CHARTS,
    "xl/media/image1.png": L.IMAGES,
    "xl/drawings/drawing1.xml": L.DRAWINGS,
    "xl/pivotCache/pivotCacheDefinition1.xml": L.PIVOT_CACHE,
    "xl/externalLinks/externalLink1.xml": L.EXTERNAL_LINKS,
    "xl/embeddings/oleObject1.bin": L.EMBEDDED_OBJECTS,
    "xl/vbaProject.bin": L.MACROS,
}


@pytest.mark.parametrize("name", sorted(FEATURES))
def test_unextracted_package_content_is_partial_with_limitation(name) -> None:
    result = parse_spreadsheet(rezip(simple(), extra={name: b"x"}))
    check(result.outcome is Outcome.INCOMPLETE, "partial")
    check(result.limitations == (FEATURES[name],), "exact limitation")
    check(len(result.cells) == 5, "values still extracted")


def test_macro_workbook_is_xlsm_and_macros_are_not_run() -> None:
    def edit(name: str, body: bytes) -> bytes:
        if name != "[Content_Types].xml":
            return body
        old = b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
        return body.replace(
            old, b"application/vnd.ms-excel.sheet.macroEnabled.main+xml"
        )

    data = rezip(simple(), edit, {"xl/vbaProject.bin": b"not-run"})
    result = parse_spreadsheet(data)
    check(result.format == "xlsm", "xlsm")
    check(result.limitations == (L.MACROS,), "macros flagged")
    check(len(result.cells) == 5, "values read")


def package(files: dict[str, bytes]) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        for name, body in files.items():
            zf.writestr(name, body)
    return out.getvalue()


@pytest.mark.parametrize(
    "data",
    [
        package(
            {
                "mimetype": b"application/vnd.oasis.opendocument.spreadsheet",
                "content.xml": b"<x/>",
            }
        ),
        package({"xl/workbook.bin": b"\x00", "[Content_Types].xml": b"<Types/>"}),
        package({"word/document.xml": b"<w/>"}),
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64,
        b"just text",
    ],
)
def test_unsupported_formats_fail_with_limitation(data) -> None:
    result = parse_spreadsheet(data)
    check(result.outcome is Outcome.FAILED and not result.cells, "failed")
    check(result.limitations == (L.UNSUPPORTED_FORMAT,), "limitation")
    check(result.to_receipt("1") is None, "no receipt for a failure")


def test_password_protected_ole_container_fails() -> None:
    marker = "EncryptedPackage".encode("utf-16-le")
    data = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32 + marker
    result = parse_spreadsheet(data)
    check(result.outcome is Outcome.FAILED, "failed")
    check(result.limitations == (L.PASSWORD_PROTECTED,), "password")


def test_zip_level_encryption_flag_fails() -> None:
    data = bytearray(simple())
    # Set the encryption bit in every local and central header flag field.
    for sig, off in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        pos = data.find(sig)
        while pos >= 0:
            data[pos + off] |= 1
            pos = data.find(sig, pos + 4)
    result = parse_spreadsheet(bytes(data))
    check(result.outcome is Outcome.FAILED, "failed")
    check(result.limitations == (L.PASSWORD_PROTECTED,), "password")


def test_empty_corrupt_and_truncated_input_never_complete() -> None:
    good = simple()
    cases = {
        b"": L.EMPTY_FILE,
        good[: len(good) // 2]: L.CORRUPT,
        b"PK\x03\x04" + b"\x00" * 40: L.CORRUPT,
    }
    for data, limitation in cases.items():
        result = parse_spreadsheet(data)
        check(result.outcome is Outcome.FAILED, "failed")
        check(result.limitations == (limitation,), "limitation")


def test_corrupt_sheet_xml_is_corrupt() -> None:
    data = rezip(simple(), lambda n, b: b"<broken" if n.endswith("sheet1.xml") else b)
    result = parse_spreadsheet(data)
    check(result.outcome is Outcome.FAILED, "failed")
    check(result.limitations == (L.CORRUPT,), "corrupt")


def test_non_bytes_argument_raises_fixed_message() -> None:
    with pytest.raises(A2ContractError, match="workbook must be bytes"):
        parse_spreadsheet("text")  # type: ignore[arg-type]


def test_repr_and_limitation_text_hide_content() -> None:
    result = parse_spreadsheet(simple())
    shown = repr(result) + repr(result.cells[0]) + repr(result.sheets[0])
    check(SECRET not in shown and "Data" not in shown, "repr leaks content")
    for item in L:
        check(item.value.replace("_", "").isalpha(), "limitation is fixed text")


def test_receipt_matches_outcome() -> None:
    complete = parse_spreadsheet(simple()).to_receipt("1")
    partial = parse_spreadsheet(formula_book()).to_receipt("1")
    check(complete is not None and complete.outcome == "complete", "complete")
    check(partial is not None and partial.outcome == "partial", "partial")
    check(partial.limit_hit == "bytes", "limit hit")  # type: ignore[union-attr]


def test_same_bytes_give_equal_results() -> None:
    data = simple()
    check(parse_spreadsheet(data) == parse_spreadsheet(data), "deterministic")


# --- hard limits -----------------------------------------------------------


def many_cells(count: int) -> bytes:
    wb = openpyxl.Workbook()
    for row in range(1, count + 1):
        active(wb).cell(row=row, column=1, value=f"v{row}")
    return save(wb)


@pytest.mark.parametrize(
    ("name", "value", "build", "limitation"),
    [
        ("MAX_CELLS", 3, lambda: many_cells(10), L.CELL_LIMIT),
        ("MAX_TEXT_CHARS", 8, lambda: many_cells(10), L.TEXT_LIMIT),
        ("MAX_CELL_CHARS", 4, simple, L.TEXT_LIMIT),
        ("MAX_BYTES", 100, simple, L.SIZE_LIMIT),
        ("MAX_UNCOMPRESSED", 100, simple, L.ZIP_LIMIT),
        ("MAX_ENTRIES", 2, simple, L.ZIP_LIMIT),
        ("MAX_GRID_AREA", 4, lambda: many_cells(10), L.DIMENSION_LIMIT),
    ],
)
def test_hard_limit_gives_partial_with_limitation(
    monkeypatch, name, value, build, limitation
) -> None:
    monkeypatch.setattr(sp, name, value)
    result = parse_spreadsheet(build())
    check(result.outcome is Outcome.INCOMPLETE, f"{name}: partial")
    check(limitation in result.limitations, f"{name}: limitation")
    check(result.to_receipt("1") is not None, "receipt exists")


def test_sheet_limit_stops_after_max_sheets(monkeypatch) -> None:
    monkeypatch.setattr(sp, "MAX_SHEETS", 2)
    wb = openpyxl.Workbook()
    for index in range(4):
        wb.create_sheet(f"S{index}")["A1"] = index
    result = parse_spreadsheet(save(wb))
    check(len(result.sheets) == 2 and L.SHEET_LIMIT in result.limitations, "sheets")
    check(result.outcome is Outcome.INCOMPLETE, "partial")
