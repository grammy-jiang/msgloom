"""Manager-reproduced regressions for Word parser boundary repairs."""

from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from msgloom.preparation import (
    DocumentFormat,
    DocumentLocation,
    ParserConfig,
    ParserIdentity,
    ParserLimits,
    ParserRequest,
    SavedByteReference,
    TableBlock,
    TextBlock,
)
from msgloom.preparation.parsers.word import (
    BACKEND,
    PARSER_NAME,
    PARSER_VERSION,
    parse,
)


def _request(content: bytes) -> ParserRequest:
    return ParserRequest(
        source=SavedByteReference(
            reference="blob:sha256:synthetic-word-r2",
            sha256=sha256(content).hexdigest(),
            byte_count=len(content),
        ),
        detected_format=DocumentFormat.DOCX,
        parser=ParserIdentity(PARSER_NAME, PARSER_VERSION, BACKEND),
        config=ParserConfig(profile="word-native-v1"),
        limits=ParserLimits(
            wall_time_seconds=5,
            memory_bytes=128 * 1024 * 1024,
            decompressed_bytes=4 * 1024 * 1024,
            output_bytes=1024 * 1024,
            container_members=256,
        ),
    )


def _simple_docx() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Body")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    document.save(stream)
    return stream.getvalue()


def _rewrite(
    content: bytes,
    *,
    document_xml: bytes | None = None,
    extras: dict[str, bytes] | None = None,
) -> bytes:
    output = BytesIO()
    with (
        ZipFile(BytesIO(content)) as source,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml" and document_xml is not None:
                payload = document_xml
            target.writestr(info.filename, payload)
        for name, payload in (extras or {}).items():
            target.writestr(name, payload)
    return output.getvalue()


@pytest.mark.parametrize(
    "parser",
    [
        ParserIdentity("unrelated-parser", PARSER_VERSION, BACKEND),
        ParserIdentity(PARSER_NAME, "999", BACKEND),
        ParserIdentity(PARSER_NAME, PARSER_VERSION, "false-backend"),
    ],
)
def test_rejects_mismatched_identity_before_native_parsing(
    parser: ParserIdentity,
) -> None:
    content = b"not-a-docx"
    request = replace(_request(content), parser=parser)
    with pytest.raises(ValueError, match="identity"):
        parse(request, content)


@pytest.mark.parametrize(
    "config",
    [
        ParserConfig(profile="unknown-word-profile"),
        ParserConfig(profile="word-native-v1", settings=(("unknown", "1"),)),
    ],
)
def test_rejects_unknown_profile_or_settings(config: ParserConfig) -> None:
    content = _simple_docx()
    with pytest.raises(ValueError, match="profile"):
        parse(replace(_request(content), config=config), content)


def test_doc_rejects_arbitrary_profile_before_unsupported_result() -> None:
    content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1synthetic"
    request = replace(
        _request(content),
        detected_format=DocumentFormat.DOC,
        config=ParserConfig(profile="arbitrary-conversion-profile"),
    )
    with pytest.raises(ValueError, match="profile"):
        parse(request, content)


def test_preserves_grid_before_after_coordinates_and_cell_ownership() -> None:
    stream = BytesIO()
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row in range(3):
        for column in range(3):
            table.cell(row, column).text = f"R{row + 1}C{column + 1}"

    leading = table._tbl.tr_lst[1]  # pyright: ignore[reportPrivateUsage]
    leading.remove(leading.tc_lst[0])
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), "1")
    leading.get_or_add_trPr().append(grid_before)

    trailing = table._tbl.tr_lst[2]  # pyright: ignore[reportPrivateUsage]
    trailing.remove(trailing.tc_lst[-1])
    grid_after = OxmlElement("w:gridAfter")
    grid_after.set(qn("w:val"), "1")
    trailing.get_or_add_trPr().append(grid_after)
    document.save(stream)

    content = stream.getvalue()
    output = parse(_request(content), content)
    table_block = next(
        block for block in output.blocks if isinstance(block, TableBlock)
    )
    cells = {
        (cell.row_index, cell.column_index): cell.text
        for cell in table_block.table.cells
    }
    if cells.get((2, 2)) != "R2C2" or cells.get((2, 3)) != "R2C3":
        pytest.fail("Leading grid omission shifted source cell ownership")
    if cells.get((3, 1)) != "R3C1" or cells.get((3, 2)) != "R3C2":
        pytest.fail("Trailing grid omission shifted subsequent row coordinates")
    if (2, 1) in cells or (3, 3) in cells:
        pytest.fail("Omitted grid positions must not become invented cells")


def test_vertical_merge_respects_grid_before_on_continuation_rows() -> None:
    stream = BytesIO()
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 1).text = "Vertical"
    table.cell(0, 1).merge(table.cell(2, 1))
    for row_index in (1, 2):
        tr = table._tbl.tr_lst[row_index]  # pyright: ignore[reportPrivateUsage]
        tr.remove(tr.tc_lst[0])
        grid_before = OxmlElement("w:gridBefore")
        grid_before.set(qn("w:val"), "1")
        tr.get_or_add_trPr().append(grid_before)
    document.save(stream)

    content = stream.getvalue()
    output = parse(_request(content), content)
    table_block = next(
        block for block in output.blocks if isinstance(block, TableBlock)
    )
    merged = [
        cell for cell in table_block.table.cells if cell.merged_range == "R1C2:R3C2"
    ]
    if len(merged) != 1 or merged[0].text != "Vertical":
        pytest.fail("Expected vertical merge ownership at logical column two")


def test_malformed_grid_fails_instead_of_misattributing() -> None:
    stream = BytesIO()
    document = Document()
    table = document.add_table(rows=2, cols=3)
    tr = table._tbl.tr_lst[1]  # pyright: ignore[reportPrivateUsage]
    tr.remove(tr.tc_lst[0])
    document.save(stream)
    content = stream.getvalue()

    with pytest.raises(ValueError, match="grid/merge"):
        parse(_request(content), content)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_rejects_dtd_independent_of_xml_encoding(encoding: str) -> None:
    xml = '<?xml version="1.0"?><!DOCTYPE root []><root/>'.encode(encoding)
    content = _rewrite(
        _simple_docx(),
        extras={"word/synthetic.xml": xml},
    )
    with pytest.raises(ValueError, match="DTD or entity"):
        parse(_request(content), content)


def test_comments_are_safe_and_do_not_hide_feature_scans() -> None:
    base = _simple_docx()
    with ZipFile(BytesIO(base)) as archive:
        document_xml = archive.read("word/document.xml")
    marker = b"<w:sectPr"
    injected = (
        b"<!-- ordinary comment -->"
        b'<w:p><w:ins w:id="8" w:author="Synthetic">'
        b"<w:r><w:t>Revision</w:t></w:r></w:ins></w:p>"
    )
    content = _rewrite(
        base,
        document_xml=document_xml.replace(marker, injected + marker, 1),
    )
    output = parse(_request(content), content)

    codes = {limitation.code for limitation in output.limitations}
    if "revisions" not in codes:
        pytest.fail("Comment skipping must not suppress feature scans")
    located = {
        block.location.part: block.text
        for block in output.blocks
        if isinstance(block, TextBlock)
        and isinstance(block.location, DocumentLocation)
        and block.text
    }
    if located.get("word/header1.xml") != "Header":
        pytest.fail("Comment handling disrupted header extraction")
    if located.get("word/footer1.xml") != "Footer":
        pytest.fail("Comment handling disrupted footer extraction")
