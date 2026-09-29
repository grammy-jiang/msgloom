"""Focused synthetic tests for the bounded Word parser lane."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from msgloom.preparation import (
    DocumentFormat,
    DocumentLocation,
    LimitationKind,
    ParserConfig,
    ParserIdentity,
    ParserLimits,
    ParserRequest,
    SavedByteReference,
    TableBlock,
    TextBlock,
)
from msgloom.preparation.parsers.word import (
    BACKEND,
    PARSER_NAME,
    PARSER_VERSION,
    parse,
)


def _add_hyperlink(paragraph: Paragraph, text: str, target: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _base_docx() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    table.cell(1, 0).merge(table.cell(1, 1))
    paragraph = document.add_paragraph("After ")
    _add_hyperlink(paragraph, "reference", "https://example.invalid/word")
    paragraph.add_run(" tail")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header text"
    section.footer.paragraphs[0].text = "Footer text"
    document.save(stream)
    return stream.getvalue()


def _rewrite_docx(
    content: bytes,
    *,
    document_transform: Callable[[bytes], bytes] | None = None,
    member_transforms: dict[str, Callable[[bytes], bytes]] | None = None,
    extras: dict[str, bytes] | None = None,
) -> bytes:
    source = BytesIO(content)
    output = BytesIO()
    with ZipFile(source) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename == "word/document.xml" and document_transform is not None:
                payload = document_transform(payload)
            transform = (member_transforms or {}).get(info.filename)
            if transform is not None:
                payload = transform(payload)
            rewritten.writestr(info.filename, payload)
        for name, payload in (extras or {}).items():
            rewritten.writestr(name, payload)
    return output.getvalue()


def _mixed_docx() -> bytes:
    def transform(xml: bytes) -> bytes:
        marker = b"<w:sectPr"
        insert = (
            b'<w:p><w:ins w:id="7" w:author="Synthetic">'
            b"<w:r><w:t>Inserted revision</w:t></w:r></w:ins></w:p>"
            b"<w:p><w:r><w:object><o:OLEObject xmlns:o="
            b'"urn:schemas-microsoft-com:office:office" '
            b'Type="Embed" ProgID="Synthetic.Object"/></w:object></w:r></w:p>'
        )
        return xml.replace(marker, insert + marker, 1)

    return _rewrite_docx(_base_docx(), document_transform=transform)


def _mark_first_member_encrypted(content: bytes) -> bytes:
    data = bytearray(content)
    local = data.find(b"PK\x03\x04")
    central = data.find(b"PK\x01\x02")
    if local < 0 or central < 0:
        raise RuntimeError("synthetic DOCX lacks ZIP headers")
    local_flags = int.from_bytes(data[local + 6 : local + 8], "little") | 1
    central_flags = int.from_bytes(data[central + 8 : central + 10], "little") | 1
    data[local + 6 : local + 8] = local_flags.to_bytes(2, "little")
    data[central + 8 : central + 10] = central_flags.to_bytes(2, "little")
    return bytes(data)


def _request(
    content: bytes,
    *,
    detected_format: DocumentFormat = DocumentFormat.DOCX,
    decompressed_bytes: int = 4 * 1024 * 1024,
    output_bytes: int = 1024 * 1024,
    container_members: int = 256,
) -> ParserRequest:
    return ParserRequest(
        source=SavedByteReference(
            reference="blob:sha256:synthetic-word",
            sha256=sha256(content).hexdigest(),
            byte_count=len(content),
        ),
        detected_format=detected_format,
        parser=ParserIdentity(
            name=PARSER_NAME,
            version=PARSER_VERSION,
            backend=BACKEND,
        ),
        config=ParserConfig(profile="word-native-v1"),
        limits=ParserLimits(
            wall_time_seconds=5,
            memory_bytes=128 * 1024 * 1024,
            decompressed_bytes=decompressed_bytes,
            output_bytes=output_bytes,
            container_members=container_members,
        ),
    )


def test_docx_preserves_body_order_table_merge_header_footer_and_link() -> None:
    """Supported Word structure retains deterministic order and locations."""
    content = _base_docx()
    request = _request(content)
    output = parse(request, content)

    if output.provenance.source != request.source:
        pytest.fail("Expected exact saved-byte provenance")
    body = [
        block
        for block in output.blocks
        if isinstance(block, (TextBlock, TableBlock))
        and (
            isinstance(block.location, DocumentLocation)
            and block.location.part == "word/document.xml"
            if isinstance(block, TextBlock)
            else isinstance(block.table.location, DocumentLocation)
            and block.table.location.part == "word/document.xml"
        )
    ]
    if [type(block) for block in body[:3]] != [TextBlock, TableBlock, TextBlock]:
        pytest.fail("Expected paragraph/table/paragraph body order")
    if not isinstance(body[1], TableBlock):
        pytest.fail("Expected table block in body order")
    table = body[1].table
    if (table.row_count, table.column_count) != (2, 2):
        pytest.fail("Expected exact synthetic table dimensions")
    merged = [cell for cell in table.cells if cell.merged_range is not None]
    if len(merged) != 1 or merged[0].merged_range != "R2C1:R2C2":
        pytest.fail("Expected horizontal merged span metadata")
    after = body[2]
    if not isinstance(after, TextBlock):
        pytest.fail("Expected trailing paragraph")
    if after.text != "After reference tail":
        pytest.fail(f"Unexpected hyperlink paragraph text: {after.text!r}")
    if len(after.links) != 1:
        pytest.fail("Expected hyperlink retained as data")
    if after.links[0].target != "https://example.invalid/word":
        pytest.fail("Expected inert hyperlink target")
    if not isinstance(after.links[0].location, DocumentLocation):
        pytest.fail("Expected hyperlink document location")

    located_text = {
        block.location.part: block.text
        for block in output.blocks
        if (
            isinstance(block, TextBlock)
            and isinstance(block.location, DocumentLocation)
            and block.text
        )
    }
    if located_text.get("word/header1.xml") != "Header text":
        pytest.fail("Expected header extraction with exact part mapping")
    if located_text.get("word/footer1.xml") != "Footer text":
        pytest.fail("Expected footer extraction with exact part mapping")


def test_docx_retains_vertical_merged_span() -> None:
    """Vertical merge restart/continuations become one structured cell span."""
    stream = BytesIO()
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Vertical"
    table.cell(0, 0).merge(table.cell(2, 0))
    document.save(stream)
    content = stream.getvalue()
    output = parse(_request(content), content)
    table_block = next(
        block for block in output.blocks if isinstance(block, TableBlock)
    )
    merged = [cell for cell in table_block.table.cells if cell.merged_range]
    if len(merged) != 1 or merged[0].merged_range != "R1C1:R3C1":
        pytest.fail("Expected exact vertical merged-cell span")


def test_docx_reports_detected_revisions_objects_and_coverage_limit() -> None:
    """Unsupported Word features stay visible even when text parsing succeeds."""
    content = _mixed_docx()
    output = parse(_request(content), content)
    codes = {limitation.code for limitation in output.limitations}
    for code in ("revisions", "embedded-objects", "word-feature-coverage"):
        if code not in codes:
            pytest.fail(f"Expected visible Word limitation {code}")
    if not all(
        limitation.kind in (LimitationKind.PARTIAL, LimitationKind.UNSUPPORTED)
        for limitation in output.limitations
    ):
        pytest.fail("Expected unsupported feature gaps to remain explicit")


def test_docx_discloses_textbox_notes_and_unsupported_external_relationship() -> None:
    """Detected unsupported stories and relationships remain visible gaps."""

    def add_textbox(xml: bytes) -> bytes:
        marker = b"<w:sectPr"
        textbox = (
            b"<w:p><w:r><w:txbxContent><w:p><w:r><w:t>Box text</w:t>"
            b"</w:r></w:p></w:txbxContent></w:r></w:p>"
        )
        return xml.replace(marker, textbox + marker, 1)

    def add_external_relationship(xml: bytes) -> bytes:
        relationship = (
            b'<Relationship Id="rSynthetic" '
            b'Type="https://example.invalid/relationships/custom" '
            b'Target="https://example.invalid/resource" TargetMode="External"/>'
        )
        return xml.replace(b"</Relationships>", relationship + b"</Relationships>", 1)

    content = _rewrite_docx(
        _base_docx(),
        document_transform=add_textbox,
        member_transforms={
            "word/_rels/document.xml.rels": add_external_relationship,
        },
        extras={
            "word/comments.xml": b"<comments/>",
            "word/footnotes.xml": b"<footnotes/>",
        },
    )
    output = parse(_request(content), content)
    codes = {limitation.code for limitation in output.limitations}
    expected = {"comments", "footnotes", "text-boxes", "unsupported-relationships"}
    if not expected.issubset(codes):
        pytest.fail(f"Missing unsupported Word feature limitations: {expected - codes}")


def test_doc_is_explicitly_unsupported_without_conversion_profile() -> None:
    """Legacy DOC never triggers conversion or masquerades as parsed DOCX."""
    content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1synthetic"
    output = parse(_request(content, detected_format=DocumentFormat.DOC), content)
    if output.blocks:
        pytest.fail("Unsupported DOC must not produce invented content")
    if [item.code for item in output.limitations] != ["legacy-doc-unsupported"]:
        pytest.fail("Expected the explicit no-conversion DOC limitation")


@pytest.mark.parametrize(
    ("mutator", "match"),
    [
        (lambda data: b"not-a-zip", "DOCX ZIP container"),
        (
            lambda data: _rewrite_docx(
                data,
                document_transform=lambda xml: (
                    b'<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
                    + xml
                ),
            ),
            "DTD or entity",
        ),
        (
            lambda data: _rewrite_docx(
                data,
                extras={"word/vbaProject.bin": b"synthetic macro bytes"},
            ),
            "macro",
        ),
        (_mark_first_member_encrypted, "encrypted"),
        (
            lambda data: _rewrite_docx(
                data,
                document_transform=lambda xml: b"<broken",
            ),
            "malformed DOCX XML",
        ),
    ],
)
def test_docx_rejects_unsafe_or_mismatched_containers(
    mutator: Callable[[bytes], bytes], match: str
) -> None:
    """Unsafe XML/package features fail before python-docx parses content."""
    content = mutator(_base_docx())
    with pytest.raises(ValueError, match=match):
        parse(_request(content), content)


def test_docx_enforces_member_and_decompression_limits() -> None:
    """Container inventory is bounded before XML/document traversal."""
    content = _base_docx()
    with pytest.raises(ValueError, match="container member limit"):
        parse(_request(content, container_members=1), content)
    with pytest.raises(ValueError, match="decompressed byte limit"):
        parse(_request(content, decompressed_bytes=64), content)


def test_docx_enforces_extracted_output_limit() -> None:
    """Extracted semantic output cannot exceed the request output ceiling."""
    content = _base_docx()
    with pytest.raises(ValueError, match="output byte limit"):
        parse(_request(content, output_bytes=16), content)


def test_docx_rejects_request_byte_identity_mismatch() -> None:
    """Parser refuses bytes that do not match the already-saved evidence."""
    content = _base_docx()
    request = _request(content)
    with pytest.raises(ValueError, match="byte_count"):
        parse(replace(request, source=replace(request.source, byte_count=1)), content)
    with pytest.raises(ValueError, match="sha256"):
        parse(
            replace(request, source=replace(request.source, sha256="00" * 32)),
            content,
        )
