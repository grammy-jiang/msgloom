"""Targeted OOXML metadata pass using openpyxl with defusedxml enabled."""

from __future__ import annotations

from dataclasses import dataclass, replace
from io import BytesIO
from zipfile import ZipFile

import openpyxl.xml.functions as xml_functions
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, coordinate_to_tuple
from openpyxl.worksheet.worksheet import Worksheet

from msgloom.preparation.contracts import ParsedLink
from msgloom.preparation.parsers.excel_values import (
    cell_location,
    value_text_and_type,
)


@dataclass(frozen=True, slots=True)
class CellMetadata:
    """Metadata unavailable from python-calamine's value-oriented binding."""

    formula: str | None = None
    cached_value: str | None = None
    cached_type: str | None = None
    merged_range: str | None = None
    links: tuple[ParsedLink, ...] = ()
    text_override: str | None = None
    type_override: str | None = None
    number_format: str | None = None


@dataclass(frozen=True, slots=True)
class SheetMetadata:
    """OOXML source-cell presence, visibility, and cell metadata."""

    cells: dict[tuple[int, int], CellMetadata]
    explicit_cells: frozenset[tuple[int, int]]
    hidden_rows: frozenset[int]
    hidden_columns: frozenset[int]


def _hidden_columns(worksheet: Worksheet) -> frozenset[int]:
    hidden: set[int] = set()
    dimensions = worksheet.column_dimensions
    for key, dimension in dimensions.items():
        if not dimension.hidden:
            continue
        start = dimension.min or column_index_from_string(key)
        end = dimension.max or start
        hidden.update(range(start, end + 1))
    return frozenset(hidden)


def _base_cell_metadata(
    worksheet: Worksheet,
    explicit_cells: frozenset[tuple[int, int]],
) -> dict[tuple[int, int], CellMetadata]:
    cells: dict[tuple[int, int], CellMetadata] = {}
    for cell_row, cell_column in sorted(explicit_cells):
        cell = worksheet.cell(row=cell_row, column=cell_column)
        key = (cell_row, cell_column)
        metadata = CellMetadata()
        if cell.data_type == "e":
            metadata = replace(
                metadata,
                text_override=str(cell.value),
                type_override="error",
            )
        elif cell.data_type == "inlineStr" and cell.value is None:
            metadata = replace(
                metadata,
                text_override="",
                type_override="string",
            )
        if cell.number_format != "General":
            metadata = replace(
                metadata,
                number_format=cell.number_format,
            )
        hyperlink = cell.hyperlink
        if hyperlink is not None:
            target = hyperlink.target or hyperlink.location
            if target:
                link = ParsedLink(
                    target=str(target),
                    text=(None if cell.value is None else str(cell.value)),
                    location=cell_location(
                        worksheet.title,
                        cell_row,
                        cell_column,
                    ),
                )
                metadata = replace(metadata, links=(link,))
        if metadata != CellMetadata():
            cells[key] = metadata
    for merged in worksheet.merged_cells.ranges:
        range_ref = str(merged)
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                key = (row, column)
                cells[key] = replace(
                    cells.get(key, CellMetadata()),
                    merged_range=range_ref,
                )
    return cells


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@dataclass(frozen=True, slots=True)
class _RawCell:
    """Bounded facts openpyxl does not retain for empty cached values."""

    cache_present: bool
    cache_text: str | None
    type_code: str | None


def _worksheet_paths(archive: ZipFile) -> dict[str, str]:
    workbook = fromstring(archive.read("xl/workbook.xml"))
    relationships = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        item.get("Id"): item.get("Target")
        for item in relationships
        if item.get("Id") and item.get("Target")
    }
    relation_key = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    result: dict[str, str] = {}
    for sheet in workbook.iter(f"{{{_MAIN_NS}}}sheet"):
        name = sheet.get("name")
        relation = sheet.get(relation_key)
        target = rel_targets.get(relation)
        if name is None or target is None:
            raise ValueError("OOXML worksheet relationship metadata is incomplete")
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = "xl/" + target.lstrip("/")
        result[name] = path
    return result


def _raw_cells(
    archive: ZipFile,
    path: str,
) -> dict[tuple[int, int], _RawCell]:
    root = fromstring(archive.read(path))
    result: dict[tuple[int, int], _RawCell] = {}
    for element in root.iter(f"{{{_MAIN_NS}}}c"):
        coordinate = element.get("r")
        if not coordinate:
            raise ValueError("OOXML cell is missing a source coordinate")
        row, column = coordinate_to_tuple(coordinate)
        value = element.find(f"{{{_MAIN_NS}}}v")
        result[(row, column)] = _RawCell(
            cache_present=value is not None,
            cache_text=None if value is None else (value.text or ""),
            type_code=element.get("t"),
        )
    return result


def _cached_result(
    cached_cell: object,
    raw: _RawCell,
) -> tuple[str | None, str | None]:
    if not raw.cache_present:
        return None, None
    data_type = getattr(cached_cell, "data_type", None)
    value = getattr(cached_cell, "value", None)
    if data_type == "e":
        return raw.cache_text or str(value), "error"
    if data_type == "b":
        return ("true" if bool(value) else "false"), "bool"
    if value is None:
        if raw.type_code in {"str", "inlineStr"}:
            return "", "string"
        if raw.cache_text == "":
            return None, None
        return raw.cache_text, "numeric"
    text, value_type = value_text_and_type(value)
    return text, value_type


def read_ooxml_metadata(content: bytes) -> dict[str, SheetMetadata]:
    """Read formulas, caches, merges, links, formats, and visibility only."""
    if vars(xml_functions).get("DEFUSEDXML") is not True:
        raise RuntimeError("openpyxl metadata pass requires defusedxml")
    formulas = None
    cached = None
    try:
        formulas = load_workbook(
            BytesIO(content),
            read_only=False,
            keep_vba=False,
            data_only=False,
            keep_links=False,
        )
        cached = load_workbook(
            BytesIO(content),
            read_only=False,
            keep_vba=False,
            data_only=True,
            keep_links=False,
        )
        result: dict[str, SheetMetadata] = {}
        with ZipFile(BytesIO(content)) as archive:
            worksheet_paths = _worksheet_paths(archive)
            for worksheet in formulas.worksheets:
                path = worksheet_paths.get(worksheet.title)
                if path is None:
                    raise ValueError("OOXML worksheet path metadata is missing")
                raw_cells = _raw_cells(archive, path)
                cells = _base_cell_metadata(
                    worksheet,
                    frozenset(raw_cells),
                )
                cached_sheet = cached[worksheet.title]
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        cell_row = cell.row
                        cell_column = cell.column
                        if cell_row is None or cell_column is None:
                            raise ValueError(
                                "OOXML formula cell is missing a source coordinate"
                            )
                        key = (cell_row, cell_column)
                        raw = raw_cells.get(key)
                        if raw is None:
                            raise ValueError(
                                "OOXML formula is missing source cell metadata"
                            )
                        cached_cell = cached_sheet.cell(
                            row=cell_row,
                            column=cell_column,
                        )
                        cached_text, cached_type = _cached_result(
                            cached_cell,
                            raw,
                        )
                        cells[key] = replace(
                            cells.get(key, CellMetadata()),
                            formula=str(cell.value),
                            cached_value=cached_text,
                            cached_type=cached_type,
                        )
                hidden_rows = frozenset(
                    index
                    for index, dimension in worksheet.row_dimensions.items()
                    if dimension.hidden
                )
                result[worksheet.title] = SheetMetadata(
                    cells=cells,
                    explicit_cells=frozenset(raw_cells),
                    hidden_rows=hidden_rows,
                    hidden_columns=_hidden_columns(worksheet),
                )
        return result
    finally:
        if cached is not None:
            cached.close()
        if formulas is not None:
            formulas.close()
