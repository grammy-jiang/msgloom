"""Bounded bytes-only spreadsheet extraction for an isolated parser worker."""

from __future__ import annotations

from io import BytesIO

from openpyxl.utils import get_column_letter
from python_calamine import (
    CalamineError,
    CalamineSheet,
    CalamineWorkbook,
    PasswordError,
    SheetTypeEnum,
    SheetVisibleEnum,
)

from msgloom.preparation.contracts import (
    DocumentFormat,
    LimitationKind,
    ParsedCell,
    ParsedTable,
    ParserIdentity,
    ParserLimitation,
    ParserOutput,
    ParserProvenance,
    ParserRequest,
    SheetLocation,
    TableBlock,
)
from msgloom.preparation.parsers.excel_container import preflight_container
from msgloom.preparation.parsers.excel_metadata import (
    CellMetadata,
    SheetMetadata,
    read_ooxml_metadata,
)
from msgloom.preparation.parsers.excel_policy import HiddenPolicy, hidden_policy
from msgloom.preparation.parsers.excel_values import OutputBudget, build_cell

PARSER_NAME = "msgloom.excel"
PARSER_VERSION = "1"
BACKEND = "python-calamine-0.8.2+openpyxl-3.1.5"

_FORMATS = {
    DocumentFormat.XLSX,
    DocumentFormat.XLSM,
    DocumentFormat.XLS,
    DocumentFormat.XLSB,
    DocumentFormat.ODS,
}
_OOXML_FORMATS = {DocumentFormat.XLSX, DocumentFormat.XLSM}


def _expected_identity() -> ParserIdentity:
    return ParserIdentity(
        name=PARSER_NAME,
        version=PARSER_VERSION,
        backend=BACKEND,
    )


def _limitation(
    kind: LimitationKind,
    code: str,
    detail: str,
    feature: str,
    *,
    location: SheetLocation | None = None,
) -> ParserLimitation:
    return ParserLimitation(
        kind=kind,
        code=code,
        detail=detail,
        feature=feature,
        location=location,
    )


def _range_ref(start: tuple[int, int], end: tuple[int, int]) -> str:
    start_row, start_col = start
    end_row, end_col = end
    return (
        f"{get_column_letter(start_col + 1)}{start_row + 1}:"
        f"{get_column_letter(end_col + 1)}{end_row + 1}"
    )


def _calamine_merges(sheet: CalamineSheet) -> dict[tuple[int, int], str]:
    result: dict[tuple[int, int], str] = {}
    ranges = sheet.merged_cell_ranges
    if ranges is None:
        return result
    for start, end in ranges:
        ref = _range_ref(start, end)
        for row in range(start[0] + 1, end[0] + 2):
            for column in range(start[1] + 1, end[1] + 2):
                result[(row, column)] = ref
    return result


def _metadata_cell(
    metadata: SheetMetadata | None,
    row: int,
    column: int,
) -> CellMetadata:
    if metadata is None:
        return CellMetadata()
    return metadata.cells.get((row, column), CellMetadata())


def _is_hidden(
    metadata: SheetMetadata | None,
    policy: HiddenPolicy,
    row: int,
    column: int,
) -> bool:
    if metadata is None:
        return not policy.include_rows or not policy.include_columns
    if not policy.include_rows and row in metadata.hidden_rows:
        return True
    return not policy.include_columns and column in metadata.hidden_columns


def _cell_from_value(
    sheet_name: str,
    row: int,
    column: int,
    value: object,
    metadata: CellMetadata,
    calamine_merge: str | None,
) -> ParsedCell:
    merged_range = metadata.merged_range or calamine_merge
    text_override = metadata.text_override
    type_override = metadata.type_override
    if metadata.formula is not None and metadata.cached_value is not None:
        text_override = metadata.cached_value
        type_override = metadata.cached_type
    return build_cell(
        sheet_name=sheet_name,
        row=row,
        column=column,
        value=value,
        formula=metadata.formula,
        cached_value=metadata.cached_value,
        merged_range=merged_range,
        links=metadata.links,
        text_override=text_override,
        type_override=type_override,
    )


def _sheet_cells(
    sheet: CalamineSheet,
    metadata: SheetMetadata | None,
    policy: HiddenPolicy,
    budget: OutputBudget,
) -> tuple[tuple[ParsedCell, ...], int, int]:
    sheet_name = sheet.name
    values = sheet.to_python()
    start = sheet.start
    end = sheet.end
    calamine_merges = _calamine_merges(sheet)
    cells: dict[tuple[int, int], ParsedCell] = {}
    if start is not None:
        start_row, start_column = start
        for row_offset, values_row in enumerate(values):
            row = start_row + row_offset + 1
            for column_offset, value in enumerate(values_row):
                column = start_column + column_offset + 1
                key = (row, column)
                if metadata is not None:
                    if (
                        key not in metadata.explicit_cells
                        and key not in calamine_merges
                    ):
                        continue
                elif value == "":
                    continue
                if _is_hidden(metadata, policy, row, column):
                    continue
                cell_metadata = _metadata_cell(
                    metadata,
                    row,
                    column,
                )
                cell = _cell_from_value(
                    sheet_name,
                    row,
                    column,
                    value,
                    cell_metadata,
                    calamine_merges.get((row, column)),
                )
                cells[(row, column)] = cell
    structural_keys = set(calamine_merges)
    if metadata is not None:
        structural_keys.update(metadata.explicit_cells)
        structural_keys.update(metadata.cells)
    for row, column in sorted(structural_keys):
        if (row, column) in cells:
            continue
        if _is_hidden(metadata, policy, row, column):
            continue
        cell = _cell_from_value(
            sheet_name,
            row,
            column,
            None,
            _metadata_cell(metadata, row, column),
            calamine_merges.get((row, column)),
        )
        cells[(row, column)] = cell
    max_row = (end[0] + 1) if end is not None else 1
    max_column = (end[1] + 1) if end is not None else 1
    for row, column in structural_keys:
        max_row = max(max_row, row)
        max_column = max(max_column, column)
    ordered = tuple(cells[key] for key in sorted(cells))
    for cell in ordered:
        budget.charge_cell(cell)
    return ordered, max_row, max_column


def _append_limitation(
    limitations: list[ParserLimitation],
    limitation: ParserLimitation,
    budget: OutputBudget,
) -> None:
    budget.charge_limitation(
        limitation.code,
        limitation.detail,
        limitation.feature,
    )
    limitations.append(limitation)


def _unavailable_visibility_limitations(
    sheet_name: str,
    policy: HiddenPolicy,
) -> tuple[ParserLimitation, ...]:
    """Disclose cells omitted because requested visibility cannot be verified."""
    result: list[ParserLimitation] = []
    location = SheetLocation(sheet_name)
    if not policy.include_rows:
        result.append(
            _limitation(
                LimitationKind.MISSING,
                "excel-row-visibility-unavailable",
                "Row visibility metadata is unavailable; cells were omitted "
                "because hidden rows are excluded by policy.",
                "hidden-rows",
                location=location,
            )
        )
    if not policy.include_columns:
        result.append(
            _limitation(
                LimitationKind.MISSING,
                "excel-column-visibility-unavailable",
                "Column visibility metadata is unavailable; cells were omitted "
                "because hidden columns are excluded by policy.",
                "hidden-columns",
                location=location,
            )
        )
    return tuple(result)


def _metadata_limitations(
    sheet_name: str,
    metadata: SheetMetadata,
    policy: HiddenPolicy,
) -> tuple[ParserLimitation, ...]:
    result: list[ParserLimitation] = []
    location = SheetLocation(sheet_name)
    if metadata.hidden_rows and not policy.include_rows:
        result.append(
            _limitation(
                LimitationKind.PARTIAL,
                "excel-hidden-rows-excluded",
                "Hidden rows were excluded by the default/requested policy: "
                + ", ".join(str(value) for value in sorted(metadata.hidden_rows)),
                "hidden-rows",
                location=location,
            )
        )
    if metadata.hidden_columns and not policy.include_columns:
        result.append(
            _limitation(
                LimitationKind.PARTIAL,
                "excel-hidden-columns-excluded",
                "Hidden columns were excluded by the default/requested policy: "
                + ", ".join(
                    get_column_letter(value)
                    for value in sorted(metadata.hidden_columns)
                ),
                "hidden-columns",
                location=location,
            )
        )
    for (row, column), cell in sorted(metadata.cells.items()):
        if _is_hidden(metadata, policy, row, column):
            continue
        if cell.formula is not None and cell.cached_value is None:
            result.append(
                _limitation(
                    LimitationKind.MISSING,
                    "excel-formula-cache-unavailable",
                    f"No cached result is stored for {sheet_name}!"
                    f"{get_column_letter(column)}{row}; the formula was not "
                    "evaluated.",
                    "formula-cache",
                    location=location,
                )
            )
        if cell.number_format is not None:
            result.append(
                _limitation(
                    LimitationKind.PARTIAL,
                    "excel-number-format-metadata",
                    f"{sheet_name}!{get_column_letter(column)}{row} uses "
                    f"number format {cell.number_format!r}; the raw typed "
                    "value is preserved without display rendering.",
                    "number-format",
                    location=location,
                )
            )
    return tuple(result)


def _parse_workbook(
    request: ParserRequest,
    content: bytes,
    metadata: dict[str, SheetMetadata],
    policy: HiddenPolicy,
    limitations: list[ParserLimitation],
    budget: OutputBudget,
) -> tuple[TableBlock, ...]:
    blocks: list[TableBlock] = []
    try:
        workbook = CalamineWorkbook.from_filelike(BytesIO(content))
    except PasswordError:
        _append_limitation(
            limitations,
            _limitation(
                LimitationKind.ENCRYPTED,
                "excel-password-protected",
                "The spreadsheet is password protected and was not parsed.",
                "encryption",
            ),
            budget,
        )
        return ()
    except CalamineError as exc:
        raise ValueError(
            f"calamine rejected spreadsheet container: {type(exc).__name__}"
        ) from exc
    try:
        sheet_metadata = {item.name: item for item in workbook.sheets_metadata}
        for sheet_index, sheet_name in enumerate(workbook.sheet_names):
            identity = sheet_metadata[sheet_name]
            location = SheetLocation(sheet_name)
            if identity.typ is not SheetTypeEnum.WorkSheet:
                _append_limitation(
                    limitations,
                    _limitation(
                        LimitationKind.UNSUPPORTED,
                        "excel-non-worksheet-skipped",
                        f"Sheet {sheet_name!r} has unsupported type "
                        f"{identity.typ.name}.",
                        "sheet-type",
                        location=location,
                    ),
                    budget,
                )
                continue
            if (
                identity.visible is not SheetVisibleEnum.Visible
                and not policy.include_sheets
            ):
                _append_limitation(
                    limitations,
                    _limitation(
                        LimitationKind.PARTIAL,
                        "excel-hidden-sheet-excluded",
                        f"Hidden sheet {sheet_name!r} was excluded by the "
                        "default/requested policy.",
                        "hidden-sheets",
                        location=location,
                    ),
                    budget,
                )
                continue
            sheet = workbook.get_sheet_by_index(sheet_index)
            sheet_extra = metadata.get(sheet_name)
            if sheet_extra is None:
                sheet_limitations = _unavailable_visibility_limitations(
                    sheet_name,
                    policy,
                )
            else:
                sheet_limitations = _metadata_limitations(
                    sheet_name,
                    sheet_extra,
                    policy,
                )
            for limitation in sheet_limitations:
                _append_limitation(limitations, limitation, budget)
            cells, row_count, column_count = _sheet_cells(
                sheet,
                sheet_extra,
                policy,
                budget,
            )
            budget.charge_block(sheet_name)
            blocks.append(
                TableBlock(
                    order=len(blocks),
                    table=ParsedTable(
                        row_count=row_count,
                        column_count=column_count,
                        cells=cells,
                        location=location,
                    ),
                )
            )
    finally:
        workbook.close()
    return tuple(blocks)


def parse(request: ParserRequest, content: bytes) -> ParserOutput:
    """Parse saved spreadsheet bytes synchronously with no external access."""
    if request.detected_format not in _FORMATS:
        raise ValueError(
            f"msgloom.excel does not accept {request.detected_format.value}"
        )
    if request.parser != _expected_identity():
        raise ValueError("Excel parser identity does not match registered backend")
    policy = hidden_policy(request)
    facts = preflight_container(request, content)
    budget = OutputBudget(request.limits.output_bytes)
    limitations: list[ParserLimitation] = []
    for limitation in facts.limitations:
        _append_limitation(limitations, limitation, budget)
    if facts.fatal_limitation is not None:
        _append_limitation(limitations, facts.fatal_limitation, budget)
        return ParserOutput(
            provenance=ParserProvenance.from_request(request),
            limitations=tuple(limitations),
        )

    metadata: dict[str, SheetMetadata] = {}
    if request.detected_format in _OOXML_FORMATS and facts.members:
        try:
            metadata = read_ooxml_metadata(content)
        except (
            OSError,
            RuntimeError,
            TypeError,
            ValueError,
            KeyError,
            IndexError,
        ) as exc:
            _append_limitation(
                limitations,
                _limitation(
                    LimitationKind.PARTIAL,
                    "excel-ooxml-metadata-failed",
                    "The targeted OOXML metadata pass failed with "
                    f"{type(exc).__name__}; calamine values remain available.",
                    "metadata",
                ),
                budget,
            )

    blocks = _parse_workbook(
        request,
        content,
        metadata,
        policy,
        limitations,
        budget,
    )
    return ParserOutput(
        provenance=ParserProvenance.from_request(request),
        blocks=blocks,
        limitations=tuple(limitations),
    )
