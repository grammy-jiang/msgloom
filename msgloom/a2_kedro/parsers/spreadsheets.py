"""
Spreadsheet parser for saved XLSX and XLSM bytes (P2-05).

The parser reads the saved bytes of one component and returns a frozen
:class:`SheetResult`. It never raises for workbook content. A bad workbook
gives an outcome and closed, content-free :class:`Limitation` codes.
Dropped content is never ``COMPLETE``.

Admitted: XLSX and XLSM through ``openpyxl`` (read-only, ``keep_vba=False``,
``keep_links=False``: no macro runs, no external link is followed). Two
passes keep both views of a formula: ``data_only=False`` gives the formula
text and ``data_only=True`` gives the cached value. Nothing is evaluated.
A formula without a cached value is flagged.

Refused: ODS and XLSB (``UNSUPPORTED_FORMAT``, ``FAILED``).
``python-calamine`` reads them, but it allocates a dense grid from the
first to the last cell. One far cell (``A1`` plus ``XFD1048576``) aborts
the process with no Python exception. A bounded pre-scan does not fit this
module (500-line rule). See ``docs/notes/a2-v2-parser-p2-05.md``.

Hostile dimensions: the declared ``<dimension>`` is never trusted. A
bounded expat scan of each sheet finds the real extents of cells with a
value or formula, hidden rows and columns, and merges. ``openpyxl``
iterates only those extents, capped by ``MAX_GRID_AREA``.

Bounds: input bytes, ZIP entries, actual (not declared) uncompressed size,
sheets, cells, text characters and grid area. A hit gives ``INCOMPLETE``.
Present but unextracted content (merges, comments, charts, images, drawings,
pivot caches, external links, embedded objects, macros) gets a limitation.
``repr`` shows no text, name or value. Details are fixed text.

CANDIDATE/RESEARCH, PROVISIONAL: see ``docs/notes/a2-v2-parser-p2-05.md``.
"""

from __future__ import annotations

import datetime as dt
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from enum import StrEnum
from io import BytesIO
from typing import Any, Literal
from xml.parsers import expat

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from msgloom.a2_kedro.contracts.codec import A2ContractError, sha256_hex
from msgloom.a2_kedro.contracts.parse import (
    MAX_CELLS,
    MAX_SHEETS,
    CellKind,
    ParserReceipt,
)

WORKER = "msgloom.a2_kedro.parsers.spreadsheets"

#: Provisional bounds. They are parser decisions, not production limits.
MAX_BYTES = 32 * 1024 * 1024
MAX_UNCOMPRESSED = 64 * 1024 * 1024
MAX_ENTRIES = 2048
MAX_TEXT_CHARS = 1_000_000
MAX_CELL_CHARS = 32_767
MAX_GRID_AREA = 2_000_000
_SMALL, _CHUNK, _MAX_COLUMN = 1024 * 1024, 64 * 1024, 16_384
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_OLE_ENCRYPTED = "EncryptedPackage".encode("utf-16-le")
_TRUE = ("1", "true")

Format = Literal["xlsx", "xlsm"]


class Outcome(StrEnum):
    """Closed result of one parse."""

    COMPLETE = "complete"
    INCOMPLETE = "incomplete"
    FAILED = "failed"


class Limitation(StrEnum):
    """Closed, content-free reasons that a result is not complete."""

    EMPTY_FILE = "empty_file"
    UNSUPPORTED_FORMAT = "unsupported_format"
    PASSWORD_PROTECTED = "password_protected"
    CORRUPT = "corrupt"
    SIZE_LIMIT = "size_limit"
    ZIP_LIMIT = "zip_limit"
    SHEET_LIMIT = "sheet_limit"
    CELL_LIMIT = "cell_limit"
    TEXT_LIMIT = "text_limit"
    DIMENSION_LIMIT = "dimension_limit"
    FORMULA_WITHOUT_CACHE = "formula_without_cache"
    MERGED_CELLS = "merged_cells"
    COMMENTS = "comments"
    CHARTS = "charts"
    IMAGES = "images"
    DRAWINGS = "drawings"
    PIVOT_CACHE = "pivot_cache"
    EXTERNAL_LINKS = "external_links"
    EMBEDDED_OBJECTS = "embedded_objects"
    MACROS = "macros"


L = Limitation
_FATAL = frozenset({L.EMPTY_FILE, L.UNSUPPORTED_FORMAT, L.PASSWORD_PROTECTED})
_FATAL |= {L.CORRUPT}
#: Lower-case package part prefix -> limitation.
_FEATURES = (
    ("xl/comments", L.COMMENTS),
    ("xl/threadedcomments/", L.COMMENTS),
    ("xl/charts/", L.CHARTS),
    ("xl/chartsheets/", L.CHARTS),
    ("xl/media/", L.IMAGES),
    ("xl/drawings/", L.DRAWINGS),
    ("xl/pivotcache", L.PIVOT_CACHE),
    ("xl/pivottables/", L.PIVOT_CACHE),
    ("xl/externallinks/", L.EXTERNAL_LINKS),
    ("xl/embeddings/", L.EMBEDDED_OBJECTS),
    ("xl/vbaproject.bin", L.MACROS),
)
_DATES = (dt.datetime, dt.date, dt.time, dt.timedelta)


@dataclass(frozen=True, slots=True)
class Cell:
    """One non-empty cell (zero-based position). ``repr`` hides content."""

    sheet_index: int
    row: int
    column: int
    kind: CellKind
    value: object = field(repr=False)
    formula: str | None = field(repr=False, default=None)
    formula_uncached: bool = False
    hidden_row: bool = False
    hidden_column: bool = False
    truncated: bool = False

    @property
    def ref(self) -> str:
        """Return the A1 reference."""
        return f"{get_column_letter(self.column + 1)}{self.row + 1}"

    @property
    def block_id(self) -> str:
        """Return the stable source block ID: sheet index and A1 ref."""
        return f"s{self.sheet_index}!{self.ref}"


@dataclass(frozen=True, slots=True)
class Sheet:
    """One sheet. ``repr`` hides the name. Hidden indexes are zero based."""

    index: int
    name: str = field(repr=False)
    state: str  # openpyxl: visible, hidden or veryHidden
    row_count: int
    column_count: int
    hidden_rows: tuple[int, ...] = field(repr=False, default=())
    hidden_columns: tuple[int, ...] = field(repr=False, default=())

    @property
    def block_id(self) -> str:
        """Return the stable source block ID of the sheet."""
        return f"s{self.index}"


@dataclass(frozen=True, slots=True)
class SheetResult:
    """Result of one parse. ``repr`` carries counts only."""

    format: Format | None
    outcome: Outcome
    limitations: tuple[Limitation, ...]
    input_bytes: int
    input_sha256: str
    sheets: tuple[Sheet, ...] = field(default=(), repr=False)
    cells: tuple[Cell, ...] = field(default=(), repr=False)

    def __repr__(self) -> str:
        counts = (len(self.limitations), len(self.sheets), len(self.cells))
        return f"SheetResult({self.format}, {self.outcome.value}, counts={counts})"

    def to_receipt(self, worker_version: str) -> ParserReceipt | None:
        """Return the parser receipt, or ``None`` for a ``FAILED`` result."""
        if self.outcome is Outcome.FAILED:
            return None
        found = bool(self.limitations)
        return ParserReceipt(
            worker=WORKER,
            worker_version=worker_version,
            outcome="partial" if found else "complete",
            limit_hit="bytes" if found else "none",
            input_bytes=self.input_bytes,
            part_count=1,
        )


# --- bounded scans of the saved bytes -------------------------------------


@dataclass(slots=True)
class _SheetScan:  # real extents (1-based counts) and flags of one sheet
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[int] = field(default_factory=set)
    max_row: int = 0
    max_col: int = 0
    merged: int = 0


class _Reject(Exception): ...  # DTD, entity, size mismatch or bad package


def _reject(*_args: object) -> None:
    raise _Reject


Hook = Callable[[str, dict[str, str]], None]


def _xml_parser(wanted: frozenset[str], hook: Hook) -> Any:
    parser = expat.ParserCreate(namespace_separator=" ")

    def start(name: str, attrs: dict[str, str]) -> None:
        if (local := name.rsplit(" ", 1)[-1]) in wanted:
            hook(local, {k.rsplit(" ", 1)[-1]: v for k, v in attrs.items()})

    parser.StartElementHandler = start
    parser.StartDoctypeDeclHandler = _reject
    parser.EntityDeclHandler = _reject
    return parser


def _int(text: str | None, default: int) -> int:
    if text is None or not text.isascii() or not text.isdigit() or len(text) > 9:
        return default
    return int(text)


def _sheet_parser(scan: _SheetScan) -> Any:
    at = {"row": 0, "col": 0, "cell_row": 0, "cell_col": 0}

    def hook(tag: str, a: dict[str, str]) -> None:
        if tag == "row":
            at["row"], at["col"] = _int(a.get("r"), at["row"] + 1), 0
            if a.get("hidden") in _TRUE:
                scan.hidden_rows.add(at["row"] - 1)
        elif tag == "c":
            at["col"] += 1
            if "r" in a:
                at["col"] = coordinate_to_tuple(a["r"])[1]
            at["cell_row"], at["cell_col"] = at["row"], at["col"]
        elif tag == "col" and a.get("hidden") in _TRUE:
            low = _int(a.get("min"), 0)
            high = min(_int(a.get("max"), 0), _MAX_COLUMN)
            scan.hidden_cols.update(range(low - 1, high))
        elif tag == "mergeCell":
            scan.merged += 1
        elif tag in ("v", "f", "is"):
            scan.max_row = max(scan.max_row, at["cell_row"])
            scan.max_col = max(scan.max_col, at["cell_col"])

    return _xml_parser(
        frozenset({"row", "c", "v", "f", "is", "col", "mergeCell"}), hook
    )


def _small(zf: zipfile.ZipFile, name: str) -> bytes:
    info = zf.getinfo(name)
    if info.file_size > _SMALL:
        raise _Reject
    with zf.open(info) as handle:
        return handle.read(_SMALL + 1)


def _detect(zf: zipfile.ZipFile) -> Format | None:
    names = set(zf.namelist())
    if "[Content_Types].xml" not in names:
        return None
    kinds: dict[str, str] = {}

    def hook(_tag: str, a: dict[str, str]) -> None:
        kinds[a.get("PartName", "").lower()] = a.get("ContentType", "").lower()

    _xml_parser(frozenset({"Override"}), hook).Parse(
        _small(zf, "[Content_Types].xml"), True
    )
    sheet = kinds.get("/xl/workbook.xml", "") if "xl/workbook.xml" in names else ""
    if sheet.endswith("macroenabled.main+xml"):
        return "xlsm"
    return "xlsx" if sheet.endswith("spreadsheetml.sheet.main+xml") else None


Scans = dict[str, _SheetScan]


def _scan_members(zf: zipfile.ZipFile, found: set[Limitation]) -> Scans:
    """Stream every member once. Count actual bytes, never declared bytes."""
    scan, total = {}, 0
    for info in zf.infolist():
        name = info.filename.lower()
        found.update(f for prefix, f in _FEATURES if name.startswith(prefix))
        if info.is_dir():
            continue
        parser = None
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            parser = _sheet_parser(scan.setdefault(name, _SheetScan()))
        seen = 0
        with zf.open(info) as handle:
            while chunk := handle.read(_CHUNK):
                seen, total = seen + len(chunk), total + len(chunk)
                if total > MAX_UNCOMPRESSED:
                    found.add(L.ZIP_LIMIT)
                    return scan
                if seen > info.file_size:
                    raise _Reject
                if parser is not None:
                    parser.Parse(chunk, False)
        if parser is not None:
            parser.Parse(b"", True)
    return scan


def _kind(value: object, error: bool) -> CellKind:
    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, _DATES):
        return "date"
    return "error" if error else "text"


class _Build:
    """Collects sheets and cells under the shared cell and text budgets."""

    def __init__(self) -> None:
        self.found: set[Limitation] = set()
        self.sheets: list[Sheet] = []
        self.cells: list[Cell] = []
        self.chars, self.index, self.full = 0, 0, False
        self.scan = _SheetScan()

    def add(
        self, row: int, col: int, value: Any, formula: str | None, error: bool
    ) -> None:
        kind = _kind(value, error)
        if formula is None and kind == "empty":
            return
        if len(self.cells) >= MAX_CELLS:
            self.found.add(L.CELL_LIMIT)
            self.full = True
            return
        long = isinstance(value, str) and len(value) > MAX_CELL_CHARS
        long = long or len(formula or "") > MAX_CELL_CHARS
        value = value[:MAX_CELL_CHARS] if isinstance(value, str) else value
        formula = formula[:MAX_CELL_CHARS] if formula is not None else None
        self.chars += (len(value) if isinstance(value, str) else 0) + len(formula or "")
        if self.chars > MAX_TEXT_CHARS or long:
            self.found.add(L.TEXT_LIMIT)
            self.full = self.full or self.chars > MAX_TEXT_CHARS
            if self.full:
                return
        uncached = formula is not None and value is None
        if uncached:
            self.found.add(L.FORMULA_WITHOUT_CACHE)
        hidden = (row in self.scan.hidden_rows, col in self.scan.hidden_cols)
        args = (self.index, row, col, kind, value, formula, uncached, *hidden, long)
        self.cells.append(Cell(*args))

    def close_sheet(self, name: str, state: str, first: int) -> None:
        mine = self.cells[first:]
        rows = max((c.row + 1 for c in mine), default=0)
        cols = max((c.column + 1 for c in mine), default=0)
        hidden = [
            tuple(sorted(x)) for x in (self.scan.hidden_rows, self.scan.hidden_cols)
        ]
        self.sheets.append(Sheet(self.index, name, state, rows, cols, *hidden))


def _read_openpyxl(raw: bytes, scan: Scans, build: _Build) -> None:
    opts: dict[str, Any] = {"read_only": True, "keep_links": False, "keep_vba": False}
    books: list[Any] = []
    try:
        for data_only in (False, True):
            books.append(
                openpyxl.load_workbook(BytesIO(raw), data_only=data_only, **opts)
            )
        for index, (wf, wc) in enumerate(
            zip(*(b.worksheets for b in books), strict=False)
        ):
            if index >= MAX_SHEETS:
                build.found.add(L.SHEET_LIMIT)
                break
            info = scan.get(str(wf._worksheet_path).lstrip("/").lower())
            if info is None:
                raise _Reject
            build.index, build.scan = index, info
            _walk_sheet(wf, wc, build)
    except Exception:  # noqa: BLE001 - library errors must not leak their text
        build.found.add(L.CORRUPT)
    finally:
        for book in books:
            book.close()


def _walk_sheet(wf: Any, wc: Any, build: _Build) -> None:
    first, rows, cols = len(build.cells), build.scan.max_row, build.scan.max_col
    if rows * cols > MAX_GRID_AREA:
        build.found.add(L.DIMENSION_LIMIT)
        rows = max(1, MAX_GRID_AREA // cols)
    if rows * cols and not build.full:
        opts = {"max_row": rows, "max_col": cols}
        both = zip(wf.iter_rows(**opts), wc.iter_rows(**opts), strict=False)
        for r, (frow, crow) in enumerate(both):
            for c, (fc, cc) in enumerate(zip(frow, crow, strict=False)):
                formula = None
                if fc.data_type == "f":
                    formula = str(getattr(fc.value, "text", fc.value))
                build.add(r, c, cc.value, formula, cc.data_type == "e")
            if build.full:
                break
    build.close_sheet(wf.title, wf.sheet_state, first)


# --- entry point ----------------------------------------------------------


def _run(raw: bytes, build: _Build) -> Format | None:
    found = build.found
    if not raw:
        found.add(L.EMPTY_FILE)
    elif len(raw) > MAX_BYTES:
        found.add(L.SIZE_LIMIT)
    elif raw.startswith(_OLE):
        encrypted = _OLE_ENCRYPTED in raw
        found.add(L.PASSWORD_PROTECTED if encrypted else L.UNSUPPORTED_FORMAT)
    elif not raw.startswith(b"PK\x03\x04"):
        found.add(L.UNSUPPORTED_FORMAT)
    else:
        return _run_zip(zipfile.ZipFile(BytesIO(raw)), raw, build)
    return None


def _run_zip(zf: zipfile.ZipFile, raw: bytes, build: _Build) -> Format | None:
    found, infos = build.found, zf.infolist()
    if len(infos) > MAX_ENTRIES:
        found.add(L.ZIP_LIMIT)
    elif len({i.filename for i in infos}) != len(infos):
        found.add(L.CORRUPT)
    elif any(i.flag_bits & 1 for i in infos):
        found.add(L.PASSWORD_PROTECTED)
    elif (fmt := _detect(zf)) is None:
        found.add(L.UNSUPPORTED_FORMAT)
    else:
        scan = _scan_members(zf, found)
        if L.ZIP_LIMIT not in found:
            if any(s.merged for s in scan.values()):
                found.add(L.MERGED_CELLS)
            _read_openpyxl(raw, scan, build)
        return fmt
    return None


def parse_spreadsheet(data: bytes) -> SheetResult:
    """
    Parse saved workbook bytes into sheets and cells.

    The format comes from the bytes, never from a name or a claim. Only a
    non-bytes argument raises :exc:`A2ContractError`.
    """
    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise A2ContractError("workbook must be bytes")
    raw, build, fmt = bytes(data), _Build(), None
    try:
        fmt = _run(raw, build)
    except Exception:  # noqa: BLE001 - library errors must not leak their text
        build.found.add(L.CORRUPT)
    found = build.found
    if found & _FATAL and not build.sheets:
        outcome = Outcome.FAILED
    else:
        outcome = Outcome.INCOMPLETE if found else Outcome.COMPLETE
    return SheetResult(
        format=fmt,
        outcome=outcome,
        limitations=tuple(item for item in Limitation if item in found),
        input_bytes=len(raw),
        input_sha256=sha256_hex(raw),
        sheets=tuple(build.sheets),
        cells=tuple(build.cells),
    )
